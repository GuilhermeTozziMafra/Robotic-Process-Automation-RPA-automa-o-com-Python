from docx import Document
from docx.shared import Pt

from openpyxl import load_workbook
import os

#Pegando o caminho do arquivo
nome_arquivo_alunos = "Alunos.xlsx"
planinhaDadosAlunos = load_workbook(nome_arquivo_alunos)


#Selecionando a sheet/planilha/aba
sheet_selecionada = planinhaDadosAlunos["Nomes"]

for linha in range(2, len(sheet_selecionada["A"]) + 1):

    arquivoWord = Document("Certificado2.docx")

    #Seleciona o estilo
    estilo = arquivoWord.styles["Normal"]

    #Pegando o nome do aluno quando passar na celula
    nomeAluno = sheet_selecionada['A%s' % linha].value

    for paragrafo in arquivoWord.paragraphs:

        if "@nome" in paragrafo.text:
            paragrafo.text = nomeAluno
            fonte = estilo.font
            fonte.name = "Calibri (Corpo)"
            fonte.size = Pt(24)

    #Pegando o caminho da pasta e configurando o nome do certificado
    caminhoCertificados = "C:\Users\guilherme.mafra\Desktop\Python\Word\Certificados" + nomeAluno + ".docx"


    #Salvando o certificado do aluno
    arquivoWord.save(caminhoCertificados)

print("Certificados gerados com sucesso")