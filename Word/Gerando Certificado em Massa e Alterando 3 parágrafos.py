from docx import Document
from docx.shared import Pt

from openpyxl import load_workbook

#Pegando o caminho do arquivo
nome_arquivo_alunos = "DadosAlunos.xlsx"
planinhaDadosAlunos = load_workbook(nome_arquivo_alunos)


#Selecionando a sheet/planilha/aba
sheet_selecionada = planinhaDadosAlunos["Nomes"]

for linha in range(2, len(sheet_selecionada["A"]) + 1):

    arquivoWord = Document("Certificado3.docx")

    #Seleciona o estilo
    estilo = arquivoWord.styles["Normal"]

    #Pegando o nome do aluno quando passar na celula
    nomeAluno = sheet_selecionada['A%s' % linha].value
    dia = sheet_selecionada['B%s' % linha].value
    mes = sheet_selecionada['C%s' % linha].value
    ano = sheet_selecionada['D%s' % linha].value
    nomeCurso = sheet_selecionada['E%s' % linha].value
    nomeInstrutor = sheet_selecionada['F%s' % linha].value


    for paragrafo in arquivoWord.paragraphs:

        if "@nome" in paragrafo.text:
            paragrafo.text = nomeAluno
            fonte = estilo.font
            fonte.name = "Calibri (Corpo)"
            fonte.size = Pt(24)

        paragrafoP1 = "Concluiu com sucesso o curso de "
        paragrafoP2 = ", como carga horária de 20 horas, promovido pela escola de Cursos Online em "
        paragrafoCompleto = f"{paragrafoP1} {nomeCurso} {paragrafoP2} {dia} de {mes} de {ano}."

        if "escola" in paragrafo.text:
            paragrafo.text = paragrafoCompleto
            fonte = estilo.font
            fonte.name = "Calibri (Corpo)"
            fonte.size = Pt(24)

        if "Instrutor" in paragrafo.text:
            paragrafo.text = nomeInstrutor + " - Instrutor"
            fonte = estilo.font
            fonte.name = "Calibri (Corpo)"
            fonte.size = Pt(24)

    #Pegando o caminho da pasta e configurando o nome do certificado
    caminhoCertificados = "C:\\Users\guilherme.mafra\Desktop\Python\Word\Certificados" + nomeAluno + ".docx"


    #Salvando o certificado do aluno
    arquivoWord.save(caminhoCertificados)

print("Certificados gerados com sucesso")